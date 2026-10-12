from openpyxl import Workbook
from openpyxl.styles import Font
wb = Workbook(); ws = wb.active; ws.title = "Invoices"
rows = [("Invoice", "Customer", "Amount"), ("INV-5582", "Sparrow Traders", 12400),
        ("INV-5583", "Lotus Foods", 8750), ("INV-5584", "Banyan Tech", 15300)]
for r in rows: ws.append(r)
ws.append(("Total", "", "=SUM(C2:C4)"))
st = wb.create_sheet("Staff")
for r in [("Name", "Department"), ("Kavya Iyer", "Finance"), ("Rahul Bose", "Sales"), ("Anita Kulkarni", "HR")]: st.append(r)
for sh in wb:
    for row in sh.iter_rows():
        for c in row: c.font = Font(name="Arial", bold=(c.row == 1))
    for col in "ABC": sh.column_dimensions[col].width = 20
    for c in sh["C"][1:]: c.number_format = "#,##0"
wb.save("data/05_sheet.xlsx")
